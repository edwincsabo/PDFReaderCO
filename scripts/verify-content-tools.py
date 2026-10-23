"""Independent PDF/image/spreadsheet checks for the three 1.34.0 features."""
from pathlib import Path
from io import BytesIO
import base64, csv, hashlib, json, subprocess
from pypdf import PdfReader
from PIL import Image, ImageChops, ImageDraw
from openpyxl import load_workbook

version = json.loads(Path('package.json').read_text(encoding='utf-8'))['version']
if version != '1.34.0': raise RuntimeError('Este registro corresponde a 1.34.0.')
out = Path('docs/capturas') / version
out.mkdir(parents=True, exist_ok=True)
poppler = Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/Library/bin/pdftoppm.exe'
checks = []
normal = lambda text: ' '.join(text.split())

def render(file, page, label):
    target = out / label
    subprocess.run([str(poppler), '-f', str(page), '-l', str(page), '-singlefile', '-scale-to', '1400', '-png', str(file), str(target)], check=True, capture_output=True)
    return Image.open(target.with_suffix('.png')).convert('RGB')

root = Path('test-results/content-original-images')
before, after = PdfReader(root / 'entrada.pdf'), PdfReader(root / 'editado.pdf')
assert len(before.pages) == len(after.pages) == 2
assert before.metadata.title == after.metadata.title
assert dict(before.attachments) == dict(after.attachments)
assert list(after.attachments['conservar.bin']) == [bytes([1, 3, 5, 7])]
assert normal(before.pages[0].extract_text()) == normal(after.pages[0].extract_text())
assert before.get_fields()['Conservar']['/V'] == after.get_fields()['Conservar']['/V'] == 'Respuesta original'
widgets = [r.get_object() for r in after.pages[0]['/Annots'] if r.get_object().get('/Subtype') == '/Widget']
assert len(widgets) == 1 and widgets[0]['/AP']['/N'].get_data()
parent = widgets[0].get('/Parent')
assert widgets[0].get('/V', parent.get_object().get('/V') if parent else None) == 'Respuesta original'
notes = [r.get_object() for r in after.pages[0]['/Annots'] if r.get_object().get('/Subtype') == '/Text']
assert len(notes) == 1 and notes[0]['/Contents'] == 'Nota intacta'
project = json.loads((root / 'editado.pdfco').read_text(encoding='utf-8'))
assert normal(PdfReader(BytesIO(base64.b64decode(project['pdf']))).pages[0].extract_text()) == normal(after.pages[0].extract_text())
png = Image.open(root / 'extraida.png').convert('RGBA')
assert png.size == (24, 16) and png.getpixel((0, 0))[3] == 0
assert png.getpixel((16, 0)) == (20, 70, 220, 255), png.getpixel((16, 0))
a, b = render(root / 'entrada.pdf', 2, 'imagenes-pagina2-antes'), render(root / 'editado.pdf', 2, 'imagenes-pagina2-despues')
assert ImageChops.difference(a, b).getbbox() is None
a, b = render(root / 'entrada.pdf', 1, 'imagenes-antes'), render(root / 'editado.pdf', 1, 'imagenes-despues')
assert ImageChops.difference(a, b).getbbox() is not None
positions = json.loads((root / 'apariciones.json').read_text(encoding='utf-8'))
allowed = Image.new('L', a.size, 255)
draw = ImageDraw.Draw(allowed)
for items in positions.values():
    for index in [0, 2]:
        r = items[index]['rect']
        draw.rectangle((r['x']*a.width-5, r['y']*a.height-5, (r['x']+r['w'])*a.width+5, (r['y']+r['h'])*a.height+5), fill=0)
diff = ImageChops.difference(a, b)
diff.paste((0, 0, 0), mask=ImageChops.invert(allowed))
assert diff.getbbox() is None, 'Cambió contenido fuera de las dos apariciones editadas.'
project_pdf = out / 'imagen-proyecto-base.pdf'
project_pdf.write_bytes(base64.b64decode(project['pdf']))
project_render = render(project_pdf, 1, 'imagen-proyecto-base')
canvas = Image.open(root / 'canvas-dialogo.png').convert('RGB')
selected = positions['after'][2]['rect']
for label, visual in [('proyecto', project_render), ('interfaz', canvas)]:
    group_box = (int(.083333*visual.width), int(.4875*visual.height), int(.416667*visual.width), int(.6875*visual.height))
    group = visual.crop(group_box)
    ink = ImageChops.difference(group, Image.new('RGB', group.size, 'white')).convert('L').point(lambda p: 255 if p > 35 else 0).getbbox()
    assert ink is not None, 'La imagen sustituida no es visible en ' + label
    x0, y0, x1, y1 = (ink[0]+group_box[0], ink[1]+group_box[1], ink[2]+group_box[0], ink[3]+group_box[1])
    assert x0 >= selected['x']*visual.width-4 and y0 >= selected['y']*visual.height-4 and x1 <= (selected['x']+selected['w'])*visual.width+4 and y1 <= (selected['y']+selected['h'])*visual.height+4, ('Imagen fuera de la selección', label, ink, selected)
checks.append(dict(case='imagenes', canonical_fields_notes_attachments_preserved=True, png_native_pixels_and_alpha_preserved=True, other_page_pixels_identical=True, pixels_outside_changed_appearances_identical=True, replacement_visible_inside_selection_in_pdf_and_ui=True))

root = Path('test-results/content-tables')
expected = [['Código', 'Descripción', 'Importe'], ['0012', 'Café revisado', '1250,50'], ['0027', 'Piñón de acero', '-40,25'], ['0038', '=SUM(A1:A2)', '0']]
book = load_workbook(root / 'tabla.xlsx', data_only=False)
sheet = book.active
assert [list(row) for row in sheet.iter_rows(values_only=True)] == expected
assert all(c.data_type != 'f' for row in sheet for c in row)
with (root / 'tabla.csv').open(encoding='utf-8-sig', newline='') as f: rows = list(csv.reader(f, delimiter=';'))
assert rows[1] == expected[1] and rows[2][2] == "'-40,25" and rows[3][1] == "'=SUM(A1:A2)"
project = json.loads((root / 'tabla.pdfco').read_text(encoding='utf-8'))
assert project['tableDraft']['cells'] == expected
assert base64.b64decode(project['pdf']) == (root / 'entrada.pdf').read_bytes()
a, b = render(root / 'entrada.pdf', 1, 'tabla-antes'), render(root / 'sin-cambios.pdf', 1, 'tabla-despues')
assert ImageChops.difference(a, b).getbbox() is None
recovered = load_workbook('test-results/content-table-recovery/recuperada.xlsx', data_only=False)
assert recovered.active['B2'].value == 'Piñón recuperado' and recovered.active['A2'].value == '0012'
assert recovered.active['B4'].value == '=SUM(A1:A2)' and recovered.active['B4'].data_type != 'f'
checks.append(dict(case='tablas', xlsx_opens_with_expected_literal_values=True, csv_formula_prefixes_protected=True, project_draft_matches=True, source_pdf_bytes_and_render_preserved=True, forced_restart_corrections_preserved=True))

root = Path('test-results/content-auto-comparison')
report = json.loads((root / 'informe.json').read_text(encoding='utf-8'))
assert len(report['sources']['left']) == len(report['sources']['right']) == 64
assert report['sources']['left'] != report['sources']['right']
old, new = PdfReader(root / 'anterior.pdf'), PdfReader(root / 'nuevo.pdf')
assert len(old.pages) == 2 and len(new.pages) == 3
assert 'Pendiente' in old.pages[0].extract_text() and 'Aprobado' in new.pages[0].extract_text()
assert old.get_fields()['Dato']['/V'] == 'Valor antiguo' and new.get_fields()['Dato']['/V'] == 'Valor nuevo'
changes = report['pages'][0]['changes']
assert any(c.get('type') == 'text' and c['before'] == 'Pendiente' and c['after'] == 'Aprobado' for c in changes)
assert any(c['kind'] == 'page-added' for p in report['pages'] for c in p['changes'])
visual = [b for c in changes if c.get('type') == 'visual' for b in c['left']]
assert any(b['x'] < .3 < b['x']+b['w'] and b['y'] < .38 < b['y']+b['h'] for b in visual), 'Falta el cambio de color de la figura.'
assert any(b['y'] > .8 for b in visual), 'Falta el cambio de apariencia del formulario.'
a, b = render(root / 'anterior.pdf', 1, 'comparacion-anterior'), render(root / 'nuevo.pdf', 1, 'comparacion-nuevo')
assert ImageChops.difference(a, b).getbbox() is not None
with (root / 'informe.csv').open(encoding='utf-8-sig', newline='') as f: rows = list(csv.reader(f, delimiter=';'))
assert any('Pendiente' in row and 'Aprobado' in row for row in rows)
checks.append(dict(case='comparacion', text_changes_match_independent_extraction=True, form_and_color_changes_marked=True, added_page_reported=True, json_and_csv_consistent=True))
(out / 'verificacion-contenido.json').write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(checks, ensure_ascii=False, indent=2))
